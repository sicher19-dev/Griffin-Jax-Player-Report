#!/usr/bin/env python3
"""
No Brand Baseball Analytics - Biomech player page builder.

Reads one or more KinaTrax template workbooks (one sheet per game), maps every
value by pitch-type HEADER NAME (column order changes sheet to sheet), checks
the template layout row by row, matches each sheet to a real start by pitch-type
counts (MLB public game feed), applies logged corrections, and writes
data/<player-slug>.js for index.html.

Usage:
    python tools/build_biomech.py --player griffin-jax inputs/*.xlsx
    python tools/build_biomech.py --player griffin-jax --offline inputs/*.xlsx

Adding games: drop the new workbook(s) in inputs/ and re-run with ALL files.
Sessions are de-duplicated by verified game date.
"""
import argparse, json, sys, os, datetime, urllib.request, collections
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---------------------------------------------------------------- players
PLAYERS = {
    "griffin-jax": {
        "name": "Griffin Jax", "mlbam": 643377, "team": "Tampa Bay Rays",
        "teamAbbr": "TB", "throws": "R", "season": 2026,
        "note": "Starter since 5/26/26. IL 8/9 (right elbow), returned 9/2.",
    },
}

PITCH_ALIASES = {
    "changeup": "CH", "change up": "CH", "sinker": "SI", "2s": "SI",
    "4s": "FF", "fourseam": "FF", "four seam": "FF", "sweeper": "SW",
    "curveball": "CU", "curve": "CU", "slider": "SL", "cutter": "FC",
    "splitter": "FS",
}
FEED_TO_NB = {"FF": "FF", "SI": "SI", "CH": "CH", "ST": "SW", "CU": "CU",
              "KC": "CU", "SL": "SL", "FC": "FC", "FS": "FS"}

def pitch_code(v):
    if not isinstance(v, str):
        return None
    return PITCH_ALIASES.get(" ".join(v.strip().lower().split()))

# ------------------------------------------------------- metric catalog
# row: (key, label, group, unit, agg)   agg: wavg = pitch-weighted, sum = game total
# The template row text is checked against EXPECT so a changed template fails loudly.
C = collections.OrderedDict()
def m(row, key, label, group, unit, agg="wavg", expect=None):
    C[row] = dict(row=row, key=key, label=label, group=group, unit=unit, agg=agg,
                  expect=expect)

m(19, "stride", "Stride length", "release", "% height", expect="stride")
m(20, "rel_side", "Release side", "release", "in", expect="release side")
m(21, "rel_height", "Release height", "release", "in", expect="release height")
m(22, "rel_ext", "Release extension", "release", "in", expect="release extension")
m(24, "sh_rot_fc", "Shoulder rotation at foot strike", "arm", "°", expect="shoulder rotation at fc")
m(25, "sh_rot_mer", "Layback (max external rotation)", "arm", "°", expect="shoulder rotation ar mer")
m(26, "sh_rot_br", "Shoulder rotation at release", "arm", "°", expect="shoulder rotation at br")
m(27, "msrv", "Peak shoulder rotation speed", "sequence", "°/s", expect="max shoulder rotation velocity")
m(28, "srv_br", "Shoulder rotation speed at release", "arm", "°/s", expect="shoulder rotation velocity at br")
m(29, "msrv_br_t", "Shoulder peak to release", "sequence", "ms", expect="msrv to br")
m(30, "abd_fc", "Arm abduction at foot strike", "arm", "°", expect="shoulder abduction at fc")
m(31, "abd_mer", "Arm abduction at layback", "arm", "°", expect="shoulder abduction at mer")
m(32, "abd_br", "Arm abduction at release", "arm", "°", expect="shoulder abduction at br")
m(33, "plane_eff", "Shoulder plane efficiency", "arm", "%", expect="shoulder plane efficiency")
m(34, "hz_max", "Max scap load (horizontal abduction)", "arm", "°", expect="max shoulder horizontal")
m(35, "hz_max_t", "Foot strike to max scap load", "arm", "ms", expect="fc to max horizontal")
m(36, "hz_fc", "Scap load at foot strike", "arm", "°", expect="shoulder horizontal ad/abduction at fc")
m(37, "hz_mer", "Scap load at layback", "arm", "°", expect="shoulder horizontal ad/abduction at mer")
m(38, "hz_br", "Scap load at release", "arm", "°", expect="shoulder horizontal ad/abduction at br")
m(39, "sh_force", "Peak shoulder force", "load", "N", expect="max resultant shoulder force")
m(40, "sh_force_mer", "Shoulder force at layback", "load", "N", expect="resultant shoulder force at mer")
m(41, "sh_force_n", "Peak shoulder force (body weights)", "load", "×BW", expect="normalized max resultant")
m(42, "sh_force_mer_n", "Shoulder force at layback (body weights)", "load", "×BW", expect="normalized resultant shoulder force at mer")
m(43, "sh_torque", "Peak shoulder rotation torque", "load", "N·m", expect="max shoulder rotation torque")
m(44, "sh_torque_mer", "Shoulder rotation torque at layback", "load", "N·m", expect="shoulder rotation torque at mer")
m(45, "sh_torque_n", "Peak shoulder torque (per BW)", "load", "ratio", expect="normalized max shoulder rotation torque")
m(46, "sh_torque_mer_n", "Shoulder torque at layback (per BW)", "load", "ratio", expect="normalized shoulder rotation torque at mer")
m(57, "elb_flex_fc", "Elbow bend at foot strike", "arm", "°", expect="elbow flexion at fc")
m(58, "elb_flex_mer", "Elbow bend at layback", "arm", "°", expect="elbow flexion at mer")
m(59, "elb_flex_br", "Elbow bend at release", "arm", "°", expect="elbow flexion at br")
m(60, "elb_in90", "Elbow inside 90° (FS to release)", "arm", "%", expect="elbow flexion inside 90")
m(61, "meev", "Peak elbow extension speed", "sequence", "°/s", expect="max elbow extension velocity")
m(62, "meev_br_t", "Elbow peak to release", "sequence", "ms", expect="meev to br")
m(63, "varus", "Peak elbow varus torque", "load", "N·m", expect="max elbow varus torque")
m(64, "varus_mer", "Elbow varus torque at layback", "load", "N·m", expect="elbow varus torque at mer")
m(65, "varus_n", "Peak elbow varus torque (per BW)", "load", "ratio", expect="max elbow varus torque")
m(66, "varus_mer_n", "Elbow varus at layback (per BW)", "load", "ratio", expect="elbow varus torque at mer")
m(67, "pro_fc", "Forearm pronation at foot strike", "arm", "°", expect="elbow pronation/supination at fc")
m(68, "pro_mer", "Forearm pronation at layback", "arm", "°", expect="elbow pronation/supination at mer")
m(69, "pro_br", "Forearm pronation at release", "arm", "°", expect="elbow pronation/supination at br")
m(70, "gt_varus", "Game total: peak elbow varus", "gametotal", "N·m", "sum", expect="max elbow varus torque")
m(71, "gt_varus_mer", "Game total: elbow varus at layback", "gametotal", "N·m", "sum", expect="elbow varus torque at mer")
m(72, "gt_sh_force", "Game total: peak shoulder force", "gametotal", "N", "sum", expect="max resultant shoulder force")
m(73, "gt_sh_force_mer", "Game total: shoulder force at layback", "gametotal", "N", "sum", expect="resultant shoulder force at mer")
m(74, "gt_sh_torque", "Game total: peak shoulder torque", "gametotal", "N·m", "sum", expect="max shoulder rotation torque")
m(75, "gt_sh_torque_mer", "Game total: shoulder torque at layback", "gametotal", "N·m", "sum", expect="shoulder rotation torque at mer")
m(76, "tr_rot_fc", "Trunk rotation at foot strike", "trunk", "°", expect="trunk rotation at fc")
m(77, "sep_fc", "Hip-shoulder separation at foot strike", "trunk", "°", expect="trunk rotation at fc (wrt pelvis)")
m(78, "tr_rot_mer", "Trunk rotation at layback", "trunk", "°", expect="trunk rotation at mer")
m(79, "sep_mer", "Trunk vs pelvis at layback", "trunk", "°", expect="trunk rotation at mer (wrt pelvis)")
m(80, "tr_rot_br", "Trunk rotation at release", "trunk", "°", expect="trunk rotation at br")
m(81, "sep_br", "Trunk vs pelvis at release", "trunk", "°", expect="trunk rotation at br (wrt pelvis)")
m(82, "tr_flex_max", "Max forward trunk tilt", "trunk", "°", expect="max trunk flexion")
m(83, "tr_flex_max_p", "Max forward tilt vs pelvis", "trunk", "°", expect="max trunk flexion (wrt pelvis)")
m(84, "tr_flex_fc", "Forward trunk tilt at foot strike", "trunk", "°", expect="trunk flexion at fc")
m(85, "tr_flex_fc_p", "Forward tilt vs pelvis at foot strike", "trunk", "°", expect="trunk flexion at fc (wrt pelvis)")
m(86, "tr_flex_mer", "Forward trunk tilt at layback", "trunk", "°", expect="trunk flexion at mer")
m(87, "tr_flex_mer_p", "Forward tilt vs pelvis at layback", "trunk", "°", expect="trunk flexion at mer (wrt pelvis)")
m(88, "tr_flex_br", "Forward trunk tilt at release", "trunk", "°", expect="trunk flexion at br")
m(89, "tr_flex_br_p", "Forward tilt vs pelvis at release", "trunk", "°", expect="trunk flexion at br (wrt pelvis)")
m(90, "lean_max", "Max glove-side trunk lean", "trunk", "°", expect="max trunk lean")
m(91, "lean_max_p", "Max trunk lean vs pelvis", "trunk", "°", expect="max trunk lean (wrt pelvis)")
m(92, "lean_fc", "Trunk lean at foot strike", "trunk", "°", expect="trunk lean at fc")
m(93, "lean_fc_p", "Trunk lean vs pelvis at foot strike", "trunk", "°", expect="trunk lean at fc (wrt pelvis)")
m(94, "lean_mer", "Trunk lean at layback", "trunk", "°", expect="trunk lean at mer")
m(95, "lean_mer_p", "Trunk lean vs pelvis at layback", "trunk", "°", expect="trunk lean at mer (wrt pelvis)")
m(96, "lean_br", "Trunk lean at release", "trunk", "°", expect="trunk lean at br")
m(97, "lean_br_p", "Trunk lean vs pelvis at release", "trunk", "°", expect="trunk lean at br (wrt pelvis)")
m(98, "tr_acc", "Peak trunk rotation acceleration", "trunk", "°/s²", expect="max trunk rotation acceleration")
m(99, "mtrv", "Peak trunk rotation speed", "sequence", "°/s", expect="max trunk rotation velocity")
m(100, "pel_elb_t", "Pelvis peak to elbow peak", "sequence", "ms", expect="max pelvis to max elbow")
m(101, "tr_elb_t", "Trunk peak to elbow peak", "sequence", "ms", expect="max trunk to max elbow")
m(102, "pel_tilt_fc", "Pelvis tilt at foot strike", "pelvis", "°", expect="pelvis tilt at fc")
m(103, "pel_tilt_mer", "Pelvis tilt at layback", "pelvis", "°", expect="pelvis tilt at mer")
m(104, "pel_tilt_br", "Pelvis tilt at release", "pelvis", "°", expect="pelvis tilt at br")
m(105, "pel_rot_fc", "Pelvis rotation at foot strike", "pelvis", "°", expect="pelvis rotation at fc")
m(106, "pel_rot_mer", "Pelvis rotation at layback", "pelvis", "°", expect="pelvis rotation at mer")
m(107, "pel_rot_br", "Pelvis rotation at release", "pelvis", "°", expect="pelvis rotation at br")
m(108, "pel_rot_pct", "Pelvis rotation done after foot strike", "pelvis", "%", expect="pelvis rotation % after fc")
m(109, "pel_lean_fc", "Pelvis lean at foot strike", "pelvis", "°", expect="pelvis lean at fc")
m(110, "pel_lean_mer", "Pelvis lean at layback", "pelvis", "°", expect="pelvis lean at mer")
m(111, "pel_lean_br", "Pelvis lean at release", "pelvis", "°", expect="pelvis lean at br")
m(115, "th_rot_fc", "Back hip rotation at foot strike", "pelvis", "°", expect="trail hip rotation at fc")
m(116, "th_rot_mer", "Back hip rotation at layback", "pelvis", "°", expect="trail hip rotation at mer")
m(117, "th_rot_br", "Back hip rotation at release", "pelvis", "°", expect="trail hip rotation at br")
m(118, "mprv", "Peak pelvis rotation speed", "sequence", "°/s", expect="max pelvis rotation velocity")
m(119, "fc_pel_t", "Foot strike to pelvis peak", "sequence", "ms", expect="fc to max pelvis velocity")
m(120, "com_ml_fc", "Body drift side-to-side at foot strike", "com", "vel", expect="center of mass veloctiy at fc (medial")
m(121, "com_ml_mer", "Body drift side-to-side at layback", "com", "vel", expect="center of mass veloctiy at mer (medial")
m(122, "com_ml_br", "Body drift side-to-side at release", "com", "vel", expect="center of mass veloctiy at br (medial")
m(123, "com_v_fc", "Body vertical speed at foot strike", "com", "vel", expect="center of mass veloctiy at fc (vertical")
m(124, "com_v_mer", "Body vertical speed at layback", "com", "vel", expect="center of mass veloctiy at mer (vertical")
m(125, "com_v_br", "Body vertical speed at release", "com", "vel", expect="center of mass veloctiy at br (vertical")
m(126, "lh_flex_fc", "Front hip flexion at foot strike", "legs", "°", expect="lead hip flexion at fc")
m(127, "lh_flex_mer", "Front hip flexion at layback", "legs", "°", expect="lead hip flexion at mer")
m(128, "lh_flex_br", "Front hip flexion at release", "legs", "°", expect="lead hip flexion at br")
m(129, "lh_abd_fc", "Front hip ab/adduction at foot strike", "legs", "°", expect="lead hip ab/adduction at fc")
m(130, "lh_abd_mer", "Front hip ab/adduction at layback", "legs", "°", expect="lead hip ab/adduction at mer")
m(131, "lh_abd_br", "Front hip ab/adduction at release", "legs", "°", expect="lead hip ab/adduction at br")
m(132, "lh_rot_fc", "Front hip rotation at foot strike", "legs", "°", expect="lead hip rotation at fc")
m(133, "lh_rot_mer", "Front hip rotation at layback", "legs", "°", expect="lead hip rotation at mer")
m(134, "lh_rot_br", "Front hip rotation at release", "legs", "°", expect="lead hip rotation at br")
m(135, "th_flex_fc", "Back hip flexion at foot strike", "legs", "°", expect="trail hip flexion at fc")
m(136, "th_flex_mer", "Back hip flexion at layback", "legs", "°", expect="trail hip flexion at mer")
m(137, "th_flex_br", "Back hip flexion at release", "legs", "°", expect="trail hip flexion at br")
m(138, "th_abd_fc", "Back hip ab/adduction at foot strike", "legs", "°", expect="trail hip ab/adduction at fc")
m(139, "th_abd_mer", "Back hip ab/adduction at layback", "legs", "°", expect="trail hip ab/adduction at mer")
m(140, "th_abd_br", "Back hip ab/adduction at release", "legs", "°", expect="trail hip ab/adduction at br")
m(144, "lh_flex_vel", "Peak front hip flexion speed", "legs", "°/s", expect="max lead hip flexion velocity")
m(145, "fc_lh_rot_t", "Foot strike to front hip rotation", "legs", "ms", expect="fc to lead hip rotation")
m(146, "fc_th_rot_t", "Foot strike to back hip rotation", "legs", "ms", expect="fc to trail hip rotation")
m(147, "fc_th_ext_t", "Foot strike to back hip extension", "legs", "ms", expect="fc to trail hip extension")
m(151, "lk_flex_fc", "Front knee bend at foot strike", "legs", "°", expect="lead knee flexion at fc")
m(152, "lk_flex_mer", "Front knee bend at layback", "legs", "°", expect="lead knee flexion at mer")
m(153, "lk_flex_br", "Front knee bend at release", "legs", "°", expect="lead knee flexion at br")
m(154, "la_flex_fc", "Front ankle flexion at foot strike", "legs", "°", expect="lead ankle flexion at fc")
m(155, "la_flex_mer", "Front ankle flexion at layback", "legs", "°", expect="lead ankle flexion at mer")
m(156, "la_flex_br", "Front ankle flexion at release", "legs", "°", expect="lead ankle flexion at br")
m(160, "tk_flex_fc", "Back knee bend at foot strike", "legs", "°", expect="trail knee flexion at fc")
m(161, "lk_ext_vel", "Peak front knee extension speed", "legs", "°/s", expect="max knee extension velocity")
m(162, "lk_ext_t", "Release to peak knee extension", "legs", "ms", expect="br to max knee extension")
m(163, "lk_vel_fc", "Front knee speed at foot strike", "legs", "°/s", expect="lead knee flexion/extension velocity at fc")
m(164, "lk_vel_mer", "Front knee speed at layback", "legs", "°/s", expect="lead knee flexion/extension velocity at mer")
m(165, "lk_vel_br", "Front knee speed at release", "legs", "°/s", expect="lead knee flexion/extension velocity at br")
m(169, "la_ev_fc", "Front ankle roll at foot strike", "legs", "°", expect="lead ankle eversion/inversion at fc")
m(170, "la_ev_mer", "Front ankle roll at layback", "legs", "°", expect="lead ankle eversion/inversion at mer")
m(171, "la_ev_br", "Front ankle roll at release", "legs", "°", expect="lead ankle eversion/inversion at br")
m(172, "lf_angle", "Front foot angle at foot strike", "legs", "°", expect="lead foot angle at fc")
m(173, "tf_angle", "Back foot angle at foot strike", "legs", "°", expect="trail foot angle at fc")
m(174, "mhv", "Peak hand speed", "sequence", "°/s", expect="max hand resultant velocity")
m(175, "mhv_br_t", "Hand peak to release", "sequence", "ms", expect="mhv to br")

# Duplicate rows in the template that must match their primary row.
DUP_ROWS = {47: 24, 48: 25, 49: 26, 50: 30, 51: 31, 52: 32, 53: 34, 54: 36, 55: 37,
            56: 38, 112: 105, 113: 106, 114: 107, 141: 115, 142: 116, 143: 117,
            148: 126, 149: 127, 150: 128, 157: 151, 158: 152, 159: 153, 166: 154,
            167: 155, 168: 156}
# Game-total rows = per-pitch row x count
TOTAL_OF = {70: 63, 71: 64, 72: 39, 73: 40, 74: 43, 75: 44}

# Session-summary block (rows 5-16, column H)
SUMMARY = {5: "pelvis", 6: "fc_to_pelvis", 7: "trunk", 8: "pelvis_to_trunk",
           9: "elbow", 10: "trunk_to_elbow", 11: "shoulder", 12: "elbow_to_shoulder",
           13: "shoulder_to_br", 14: "hand", 15: "shoulder_to_hand", 16: "hand_to_br"}


def clean(s):
    return " ".join(str(s).lower().split()) if s is not None else ""


def parse_sheet(ws, log):
    raw_date = ws["E1"].value
    if isinstance(raw_date, (int, float)):
        raw_date = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=raw_date)
    if not isinstance(raw_date, datetime.datetime):
        log.append(f"[{ws.title}] no date in E1 - skipped")
        return None
    # template check
    for row, spec in C.items():
        got = clean(ws.cell(row, 2).value)
        if spec["expect"] and not got.startswith(spec["expect"]):
            raise SystemExit(f"[{ws.title}] template changed at row {row}: '{got}' (expected '{spec['expect']}')")
    # pitch columns by header name (row 18)
    cols = {}
    for c in range(5, 12):
        code = pitch_code(ws.cell(18, c).value)
        if code:
            cols[code] = c
    # arm-slot block: find 'ARM SLOT' in row 18, names in row 19, values in row 20
    slot = {}
    for c in range(10, 25):
        if clean(ws.cell(18, c).value) == "arm slot":
            for cc in range(c, c + 8):
                code = pitch_code(ws.cell(19, cc).value)
                if code and isinstance(ws.cell(20, cc).value, (int, float)):
                    slot[code] = ws.cell(20, cc).value
            break
    pitches, textfix = {}, []
    for code, c in cols.items():
        n = ws.cell(23, c).value
        if not isinstance(n, (int, float)) or n <= 0:
            continue
        vals = {}
        for row, spec in C.items():
            v = ws.cell(row, c).value
            if isinstance(v, str) and v.strip():
                t = v.strip().replace(" ", "")
                try:
                    num = float(t.replace(",", ".")) if t.count(",") == 1 and "." not in t else float(t)
                    textfix.append({"type": "text", "pitch": code, "key": spec["key"], "raw": v, "fixed": num,
                                    "why": "Typed as text in the sheet; read as a number."})
                    v = num
                except ValueError:
                    textfix.append({"type": "text", "pitch": code, "key": spec["key"], "raw": v, "fixed": None,
                                    "why": "Unreadable text in a number cell; left blank."})
                    v = None
            vals[spec["key"]] = v if isinstance(v, (int, float)) else None
        vals["arm_slot"] = slot.get(code)
        pitches[code] = {"n": int(n), "col": c, "v": vals}
    summary = {}
    for row, k in SUMMARY.items():
        v = ws.cell(row, 8).value
        summary[k] = v if isinstance(v, (int, float)) else None
    return {"sheet": ws.title, "sheetDate": raw_date.date().isoformat(),
            "pitches": pitches, "summary": summary, "ws": ws, "textfix": textfix}


def integrity(sess, log):
    """Duplicate-row and game-total checks. Returns list of flags."""
    ws = sess["ws"]
    flags = []
    for code, p in sess["pitches"].items():
        c = p["col"]
        for dup, prim in DUP_ROWS.items():
            a, b = ws.cell(prim, c).value, ws.cell(dup, c).value
            if isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) > 1e-6:
                flags.append({"type": "dup", "pitch": code, "key": C[prim]["key"],
                              "primary": a, "repeat": b, "rows": [prim, dup]})
        for tr, br in TOTAL_OF.items():
            tot, per = ws.cell(tr, c).value, ws.cell(br, c).value
            if isinstance(tot, (int, float)) and isinstance(per, (int, float)):
                exp = per * p["n"]
                if abs(tot - exp) > max(2, 0.01 * exp):
                    flags.append({"type": "total", "pitch": code, "key": C[tr]["key"],
                                  "sheet": tot, "computed": round(exp)})
    return flags


NOISY_KEYS = {"fc_th_rot_t", "fc_lh_rot_t", "hz_max_t"}  # event timings that legitimately swing sign


def scan_suspects(sessions_out, metrics):
    """Flag likely typos: lone sign flips and values far outside the same night's other pitches."""
    import statistics as st
    found = []
    for mm in metrics:
        k = mm["key"]
        if mm["group"] == "gametotal" or k in NOISY_KEYS:
            continue
        vals = [(s["date"], p, v["v"][k]) for s in sessions_out for p, v in s["pitches"].items()
                if v["v"].get(k) is not None]
        pos = [x for x in vals if x[2] > 0]
        neg = [x for x in vals if x[2] < 0]
        maj, mnr = (pos, neg) if len(pos) >= len(neg) else (neg, pos)
        if mnr and len(mnr) <= 2 and len(maj) >= 8:
            mags = [abs(x[2]) for x in maj]
            if st.median(mags) > 3:
                for x in mnr:
                    if min(mags) * 0.5 <= abs(x[2]) <= max(mags) * 1.5:
                        found.append({"type": "suspect", "date": x[0], "pitch": x[1], "key": k,
                                      "value": x[2], "why": "lone sign flip"})
        for s in sessions_out:
            sv = [(p, v["v"][k]) for p, v in s["pitches"].items() if v["v"].get(k) is not None]
            if len(sv) < 3:
                continue
            for i, (p, x) in enumerate(sv):
                others = [y for j, (q, y) in enumerate(sv) if j != i]
                med = st.median(others)
                spread = max(others) - min(others)
                if abs(med) > 3 and abs(x - med) > max(4 * spread, 0.25 * abs(med), 5):
                    if not any(f["date"] == s["date"] and f["pitch"] == p and f["key"] == k for f in found):
                        found.append({"type": "suspect", "date": s["date"], "pitch": p, "key": k,
                                      "value": x, "why": "far from the other pitches that night"})
    return found


def fetch_json(url):
    with urllib.request.urlopen(url, timeout=30) as r:
        return json.load(r)


def game_log(mlbam, season):
    d = fetch_json(f"https://statsapi.mlb.com/api/v1/people/{mlbam}/stats?stats=gameLog&group=pitching&season={season}")
    return d["stats"][0]["splits"]


def pitch_mix(pk, mlbam):
    f = fetch_json(f"https://statsapi.mlb.com/api/v1.1/game/{pk}/feed/live")
    mix, velo = collections.Counter(), collections.defaultdict(list)
    for p in f["liveData"]["plays"]["allPlays"]:
        if p["matchup"]["pitcher"]["id"] != mlbam:
            continue
        for e in p["playEvents"]:
            if e.get("isPitch"):
                t = FEED_TO_NB.get(e["details"].get("type", {}).get("code"), "OT")
                mix[t] += 1
                s = e.get("pitchData", {}).get("startSpeed")
                if s:
                    velo[t].append(s)
    g = f["gameData"]
    return dict(mix), {k: round(sum(v) / len(v), 1) for k, v in velo.items()}, g["venue"]["name"]


def match_game(sess, starts, player, log):
    """Match by exact/near pitch-type counts; prefer the sheet date when it fits."""
    mix = {k: p["n"] for k, p in sess["pitches"].items()}
    best = None
    for g in starts:
        diff = sum(abs(g["mix"].get(k, 0) - v) for k, v in mix.items())
        diff += sum(v for k, v in g["mix"].items() if k not in mix)
        over = any(v > g["mix"].get(k, 0) for k, v in mix.items())
        if over:
            continue  # captured more of a pitch than he threw -> impossible
        score = (diff, 0 if g["date"] == sess["sheetDate"] else 1)
        if best is None or score < best[0]:
            best = (score, g)
    return best


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--player", required=True)
    ap.add_argument("--offline", action="store_true")
    ap.add_argument("files", nargs="+")
    a = ap.parse_args()
    P = PLAYERS[a.player]
    corr_path = os.path.join(ROOT, "tools", "corrections.json")
    CORR = json.load(open(corr_path)).get(a.player, {}) if os.path.exists(corr_path) else {}
    sess_path = os.path.join(ROOT, "tools", "sessions.json")
    SESS = json.load(open(sess_path)).get(a.player, {}) if os.path.exists(sess_path) else {}
    log = []

    sessions = []
    for fp in a.files:
        wb = openpyxl.load_workbook(fp, data_only=True)
        for idx, ws in enumerate(wb.worksheets, 1):
            if clean(ws["B1"].value) != "date:":
                continue
            s = parse_sheet(ws, log)
            if s:
                s["file"] = os.path.basename(fp)
                s["sheet"] = f"tab {idx}"   # position in the workbook (tab titles carry vendor names)
                s["flags"] = integrity(s, log) + s["textfix"]
                sessions.append(s)

    starts = []
    if not a.offline:
        for sp in game_log(P["mlbam"], P["season"]):
            st = sp["stat"]
            if st.get("gamesStarted") != 1:
                continue
            pk = sp["game"]["gamePk"]
            try:
                mix, velo, venue = pitch_mix(pk, P["mlbam"])
            except Exception as e:  # noqa
                log.append(f"feed fail {pk}: {e}")
                continue
            starts.append({"date": sp["date"], "pk": pk, "home": sp["isHome"],
                           "opp": sp["opponent"]["name"], "venue": venue, "mix": mix, "velo": velo,
                           "line": {"ip": st["inningsPitched"], "h": st["hits"], "r": st["runs"],
                                    "er": st["earnedRuns"], "bb": st["baseOnBalls"],
                                    "k": st["strikeOuts"], "pc": st["numberOfPitches"]},
                           "dec": sp.get("isWin") and "W" or (sp.get("isLoss") and "L") or ""})
        cache = os.path.join(ROOT, "tools", f"{a.player}-starts.json")
        json.dump(starts, open(cache, "w"), indent=1)
    else:
        cache = os.path.join(ROOT, "tools", f"{a.player}-starts.json")
        starts = json.load(open(cache)) if os.path.exists(cache) else []

    OPP_ABBR = {"Texas Rangers": "TEX", "Boston Red Sox": "BOS", "Chicago White Sox": "CWS",
                "New York Mets": "NYM", "New York Yankees": "NYY", "Seattle Mariners": "SEA",
                "Toronto Blue Jays": "TOR", "Kansas City Royals": "KC", "Atlanta Braves": "ATL",
                "Baltimore Orioles": "BAL", "Athletics": "ATH", "Detroit Tigers": "DET",
                "Miami Marlins": "MIA", "Los Angeles Angels": "LAA", "Washington Nationals": "WSH",
                "Houston Astros": "HOU", "Minnesota Twins": "MIN"}

    out_sessions = {}
    held = []
    for s in sessions:
        mix = {k: p["n"] for k, p in s["pitches"].items()}
        rec = {"sheet": s["sheet"], "file": s["file"], "sheetDate": s["sheetDate"],
               "date": s["sheetDate"], "captured": sum(mix.values()), "mix": mix,
               "flags": s["flags"], "summary": s["summary"], "verified": False}
        tag = SESS.get(f"{s['file']}|{s['sheet']}")
        if tag and tag.get("kind") == "bullpen":
            a0, a1 = tag["range"]
            rec.update(kind="bullpen", date=f"{a0}", label=tag.get("label", "Bullpen"), range=[a0, a1],
                       verified=True, venue=None, oppAbbr="PEN", opp="Bullpen sessions")
            rec["flags"].append({"type": "bullpen", "range": [a0, a1], "why": tag.get("why", "")})
            m_ = None
        else:
            rec["kind"] = "game"
            m_ = match_game(s, starts, P, log) if starts else None
        if m_:
            (diff, _), g = m_
            if diff <= 4:
                rec.update(date=g["date"], pk=g["pk"], home=g["home"], opp=g["opp"],
                           oppAbbr=OPP_ABBR.get(g["opp"], g["opp"][:3].upper()),
                           venue=g["venue"], line=g["line"], thrown=g["mix"], velo=g["velo"],
                           verified=True, missed=diff)
                if g["date"] != s["sheetDate"]:
                    rec["flags"].append({"type": "date", "sheet": s["sheetDate"], "actual": g["date"],
                                         "why": "pitch mix matches this start exactly" if diff == 0 else f"pitch mix matches within {diff}"})
            else:
                log.append(f"[{s['sheet']}] no confident game match (best diff {diff})")
        if starts and not rec["verified"]:
            held.append({"file": rec["file"], "sheet": rec["sheet"], "sheetDate": rec["sheetDate"],
                         "captured": rec["captured"], "mix": rec["mix"],
                         "why": "No start matches this pitch mix, so it is held off the page until the game is confirmed."})
            continue
        # apply corrections
        for cfix in CORR.get(rec["date"], []):
            if cfix.get("keep"):
                rec["flags"].append({"type": "kept", **cfix})  # reviewed and left as-is
                continue
            p = s["pitches"][cfix["pitch"]]["v"]
            if p[cfix["key"]] == cfix["fixed"] and cfix["raw"] != cfix["fixed"]:
                continue  # already corrected in this copy of the sheet
            assert p[cfix["key"]] == cfix["raw"], f"correction raw mismatch {cfix} (sheet has {p[cfix['key']]})"
            p[cfix["key"]] = cfix["fixed"]
            rec["flags"].append({"type": "fix", **cfix})
        pitches = {}
        for code, p in s["pitches"].items():
            v = dict(p["v"])
            # game totals computed from per-pitch x count (sheet value kept in flags if off)
            for tr, br in TOTAL_OF.items():
                if v.get(C[br]["key"]) is not None:
                    v[C[tr]["key"]] = round(v[C[br]["key"]] * p["n"], 1)
            pitches[code] = {"n": p["n"], "v": v}
        rec["pitches"] = pitches
        # session aggregate
        agg = {}
        tot = sum(p["n"] for p in pitches.values())
        keys = [spec["key"] for spec in C.values()] + ["arm_slot"]
        for k in keys:
            spec = next((x for x in C.values() if x["key"] == k), {"agg": "wavg"})
            pairs = [(p["v"].get(k), p["n"]) for p in pitches.values() if p["v"].get(k) is not None]
            if not pairs:
                continue
            if spec["agg"] == "sum":
                agg[k] = round(sum(x for x, _ in pairs), 1)
            else:
                agg[k] = round(sum(x * n for x, n in pairs) / sum(n for _, n in pairs), 2)
        rec["all"] = agg
        # summary cross-check (sheet session averages vs pitch-weighted)
        chk = {"pelvis": "mprv", "trunk": "mtrv", "elbow": "meev", "shoulder": "msrv", "hand": "mhv"}
        for sk, mk in chk.items():
            sv = rec["summary"].get(sk)
            if sv is not None and mk in agg and abs(sv - agg[mk]) > max(15, 0.02 * sv):
                rec["flags"].append({"type": "summary", "key": mk, "sheet": sv, "computed": round(agg[mk])})
        if rec["date"] in out_sessions:
            old = out_sessions[rec["date"]]
            changed = []
            for code, pv in pitches.items():
                for k, v in pv["v"].items():
                    ov = old["pitches"].get(code, {}).get("v", {}).get(k)
                    if k.startswith("gt_") or ov == v:
                        continue
                    changed.append({"pitch": code, "key": k, "was": ov, "now": v})
            rec["flags"].insert(0, {"type": "resent", "file": rec["file"], "prev": old["file"],
                                    "prevSheetDate": old["sheetDate"], "changed": changed})
            log.append(f"duplicate session {rec['date']} ({s['sheet']}) - later file wins, {len(changed)} values differ")
        out_sessions[rec["date"]] = rec

    metrics = [{k: v for k, v in spec.items() if k != "expect"} for spec in C.values()]
    metrics.insert(4, {"row": 20, "key": "arm_slot", "label": "Arm slot", "group": "release",
                       "unit": "°", "agg": "wavg"})
    data = {"player": P, "built": datetime.date.today().isoformat(),
            "metrics": metrics,
            "sessions": [out_sessions[k] for k in sorted(out_sessions)], "held": held}
    reviewed = {(d_, c["pitch"], c["key"]) for d_, lst in CORR.items() for c in lst}
    for f in scan_suspects(data["sessions"], metrics):
        if (f["date"], f["pitch"], f["key"]) in reviewed:
            continue
        for sx in data["sessions"]:
            if sx["date"] == f["date"]:
                sx["flags"].append(f)
        log.append(f"SUSPECT (not corrected) {f}")
    os.makedirs(os.path.join(ROOT, "data"), exist_ok=True)
    out = os.path.join(ROOT, "data", f"{a.player}.js")
    with open(out, "w") as f:
        f.write("/* No Brand Baseball Analytics - biomech data. Generated by tools/build_biomech.py. Do not hand-edit. */\n")
        f.write("window.NB_BIOMECH = " + json.dumps(data, separators=(",", ":")) + ";\n")
    print(f"wrote {out}: {len(data['sessions'])} sessions")
    for s in data["sessions"]:
        print(f"  {s['date']} {s.get('oppAbbr','?'):>4} captured {s['captured']}/{s.get('line',{}).get('pc','?')} "
              f"verified={s['verified']} flags={len(s['flags'])}")
        for fl in s["flags"]:
            print("     ", fl)
    for l in log:
        print("LOG", l)


if __name__ == "__main__":
    main()
